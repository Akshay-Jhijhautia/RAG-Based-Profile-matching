"""Core file system tools for the resume assistant.

Every tool takes plain string paths so it can be called directly by an LLM.
``read_file``, ``write_file`` and ``search_in_file`` never raise; they return a
dict with a ``success`` flag and an ``error`` message on failure.
``list_files`` returns a list (per the spec) and therefore raises on an invalid
directory; callers such as the LLM dispatcher should catch those exceptions.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

TEXT_EXTENSIONS = {".txt", ".md"}
SUPPORTED_READ_EXTENSIONS = {".pdf", ".docx"} | TEXT_EXTENSIONS
BINARY_WRITE_EXTENSIONS = {".pdf", ".docx", ".doc"}
CONTEXT_CHARS = 80


def _resolve(path: str) -> Path:
    return Path(path).expanduser().resolve()


def _modified_iso(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")


def _error(filepath: str, message: str) -> dict:
    return {"success": False, "filepath": filepath, "error": message}


def _extract_pdf(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        # Many PDFs are "encrypted" with an empty user password; try that first.
        try:
            if not reader.decrypt(""):
                raise PermissionError("PDF is password-protected")
        except Exception as exc:
            raise PermissionError("PDF is password-protected") from exc
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip(), len(reader.pages)


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="latin-1")


def read_file(filepath: str) -> dict:
    """Read a resume file (PDF, DOCX, TXT or MD) and extract its text.

    Returns ``{"success": True, "filepath", "content", "metadata"}`` on success,
    or ``{"success": False, "filepath", "error"}`` on failure.
    """
    try:
        path = _resolve(filepath)
    except Exception as exc:
        return _error(filepath, f"Invalid path: {exc}")

    if not path.exists():
        return _error(str(path), "File not found")
    if path.is_dir():
        return _error(str(path), "Path is a directory, not a file")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_READ_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_READ_EXTENSIONS))
        return _error(str(path), f"Unsupported file type '{ext or '(none)'}'. Supported: {supported}")

    page_count = None
    try:
        if ext == ".pdf":
            content, page_count = _extract_pdf(path)
        elif ext == ".docx":
            content = _extract_docx(path)
        else:
            content = _extract_text(path)
    except PermissionError as exc:
        return _error(str(path), f"Permission denied: {exc}")
    except PdfReadError as exc:
        return _error(str(path), f"Could not parse PDF (corrupt or invalid): {exc}")
    except Exception as exc:  # corrupt DOCX zip, unexpected parser failures, etc.
        return _error(str(path), f"Failed to read file: {type(exc).__name__}: {exc}")

    metadata = {
        "name": path.name,
        "extension": ext,
        "size_bytes": path.stat().st_size,
        "modified": _modified_iso(path),
        "char_count": len(content),
        "word_count": len(content.split()),
    }
    if page_count is not None:
        metadata["page_count"] = page_count

    result = {"success": True, "filepath": str(path), "content": content, "metadata": metadata}
    if not content.strip():
        result["warning"] = (
            "No text could be extracted. The file may be empty or, for PDFs, scanned/image-only."
        )
    return result


def list_files(directory: str, extension: str = None) -> list:
    """List files (non-recursive) in ``directory``, optionally filtered by extension.

    ``extension`` is case-insensitive and may be given with or without the dot
    (``"pdf"``, ``".pdf"`` and ``".PDF"`` are equivalent).

    Returns a list of ``{"name", "path", "extension", "size_bytes", "modified"}``
    sorted by name.

    Raises:
        FileNotFoundError: if ``directory`` does not exist.
        NotADirectoryError: if ``directory`` is not a directory.
    """
    path = _resolve(directory)
    if not path.exists():
        raise FileNotFoundError(f"Directory not found: {path}")
    if not path.is_dir():
        raise NotADirectoryError(f"Not a directory: {path}")

    wanted = None
    if extension:
        wanted = extension.lower().strip()
        if not wanted.startswith("."):
            wanted = "." + wanted

    files = []
    for entry in sorted(path.iterdir(), key=lambda p: p.name.lower()):
        if not entry.is_file() or entry.name.startswith("."):
            continue
        ext = entry.suffix.lower()
        if wanted and ext != wanted:
            continue
        files.append(
            {
                "name": entry.name,
                "path": str(entry),
                "extension": ext,
                "size_bytes": entry.stat().st_size,
                "modified": _modified_iso(entry),
            }
        )
    return files


def write_file(filepath: str, content: str) -> dict:
    """Write text ``content`` to ``filepath``, creating parent directories as needed.

    Only plain-text formats are written; ``.pdf``/``.docx`` targets are rejected
    so that an invalid binary file is never produced.
    """
    try:
        path = _resolve(filepath)
    except Exception as exc:
        return _error(filepath, f"Invalid path: {exc}")

    if path.suffix.lower() in BINARY_WRITE_EXTENSIONS:
        return _error(
            str(path),
            f"Cannot write plain text to a '{path.suffix}' file. Use a .txt or .md extension instead.",
        )
    if path.is_dir():
        return _error(str(path), "Path is a directory, not a file")

    existed = path.exists()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        data = content.encode("utf-8")
        path.write_bytes(data)
    except Exception as exc:
        return _error(str(path), f"Failed to write file: {type(exc).__name__}: {exc}")

    return {
        "success": True,
        "filepath": str(path),
        "bytes_written": len(data),
        "created": not existed,
    }


def _snippet(text: str, start: int, end: int) -> str:
    lo = max(0, start - CONTEXT_CHARS)
    hi = min(len(text), end + CONTEXT_CHARS)
    snippet = " ".join(text[lo:hi].split())
    if lo > 0:
        snippet = "…" + snippet
    if hi < len(text):
        snippet = snippet + "…"
    return snippet


def search_in_file(filepath: str, keyword: str) -> dict:
    """Case-insensitive keyword search within a file's extracted text.

    Returns ``{"success", "filepath", "keyword", "match_count", "matches"}``
    where each match has ``line_number``, ``position``, ``match`` (the text as
    it appears in the file) and ``context`` (surrounding text).
    """
    if not keyword or not keyword.strip():
        return _error(filepath, "Keyword must be a non-empty string")

    read = read_file(filepath)
    if not read["success"]:
        return read

    text = read["content"]
    matches = []
    for m in re.finditer(re.escape(keyword.strip()), text, re.IGNORECASE):
        matches.append(
            {
                "line_number": text.count("\n", 0, m.start()) + 1,
                "position": m.start(),
                "match": m.group(0),
                "context": _snippet(text, m.start(), m.end()),
            }
        )

    return {
        "success": True,
        "filepath": read["filepath"],
        "keyword": keyword,
        "match_count": len(matches),
        "matches": matches,
    }
