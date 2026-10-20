"""Generate fictional sample resumes in PDF, DOCX and TXT formats.

The first four resumes are the Milestone 1 samples; the rest add variety in
roles, seniority, section headers ("Work Experience", "Technical Skills", ...)
and date formats ("2019 - Present", "Mar 2019 – Jun 2022") so the RAG
pipeline has realistic material to chunk, parse and rank.

Run from the project root:
    python scripts/generate_sample_resumes.py [output_dir]
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from fpdf import FPDF

# PDF output uses the core Helvetica font (latin-1 only), so PDF resumes must
# stick to plain ASCII hyphens; DOCX/TXT resumes may use en dashes.
RESUMES = {
    "resume_john_doe": {
        "format": "pdf",
        "name": "John Doe",
        "contact": "john.doe@example.com | (555) 010-1234 | Austin, TX",
        "sections": {
            "Summary": [
                "Backend software engineer with 6 years of experience building scalable web "
                "services in Python and Django.",
            ],
            "Skills": [
                "Python, Django, Django REST Framework, FastAPI, PostgreSQL, Redis, Celery, "
                "Docker, AWS (EC2, S3, Lambda), Git, CI/CD",
            ],
            "Experience": [
                "Senior Backend Engineer, Acme Payments (2021 - Present)",
                "- Led migration of a monolith to Python microservices, cutting p95 latency by 40%.",
                "- Designed a Celery-based job queue processing 2M transactions per day.",
                "Software Engineer, Brightside Labs (2018 - 2021)",
                "- Built REST APIs in Django serving 500k monthly users.",
                "- Wrote Python data pipelines feeding the analytics warehouse.",
            ],
            "Education": ["B.S. Computer Science, University of Texas at Austin, 2018"],
        },
    },
    "resume_jane_smith": {
        "format": "docx",
        "name": "Jane Smith",
        "contact": "jane.smith@example.com | (555) 010-5678 | Seattle, WA",
        "sections": {
            "Summary": [
                "Data scientist with 5 years of experience applying machine learning to "
                "customer analytics and forecasting.",
            ],
            "Skills": [
                "Python, pandas, NumPy, scikit-learn, PyTorch, SQL, Spark, Airflow, Tableau, "
                "A/B testing, statistics",
            ],
            "Experience": [
                "Data Scientist, Northwind Retail (2020 - Present)",
                "- Built a demand forecasting model in Python (PyTorch) improving accuracy by 18%.",
                "- Automated feature pipelines with Airflow and Spark.",
                "Data Analyst, Contoso Health (2019 - 2020)",
                "- Produced weekly KPI dashboards in Tableau and SQL.",
            ],
            "Education": ["M.S. Statistics, University of Washington, 2019"],
        },
    },
    "resume_alex_lee": {
        "format": "txt",
        "name": "Alex Lee",
        "contact": "alex.lee@example.com | (555) 010-9012 | Chicago, IL",
        "sections": {
            "Summary": [
                "Enterprise Java developer with 8 years of experience in financial services.",
            ],
            "Skills": [
                "Java 17, Spring Boot, Hibernate, Kafka, Oracle, Kubernetes, Maven, JUnit",
            ],
            "Experience": [
                "Lead Java Developer, Fabrikam Bank (2019 - Present)",
                "- Architected Spring Boot services for real-time trade settlement.",
                "- Introduced Kafka event streaming across 12 teams.",
                "Java Developer, Tailspin Insurance (2016 - 2019)",
                "- Maintained claims processing system built on Spring and Hibernate.",
            ],
            "Education": ["B.S. Software Engineering, Illinois Institute of Technology, 2016"],
        },
    },
    "resume_maria_garcia": {
        "format": "pdf",
        "name": "Maria Garcia",
        "contact": "maria.garcia@example.com | (555) 010-3456 | San Diego, CA",
        "sections": {
            "Summary": [
                "Frontend engineer focused on accessible, high-performance web applications.",
            ],
            "Skills": [
                "TypeScript, JavaScript, React, Next.js, Redux, CSS, Tailwind, Jest, Cypress, "
                "Figma, WCAG accessibility",
            ],
            "Experience": [
                "Frontend Engineer, Litware Media (2020 - Present)",
                "- Rebuilt the subscriber dashboard in React and TypeScript.",
                "- Raised Lighthouse accessibility score from 68 to 98.",
                "UI Developer, Adventure Works (2018 - 2020)",
                "- Built a shared component library used by 6 product teams.",
            ],
            "Education": ["B.A. Interaction Design, UC San Diego, 2018"],
        },
    },
    "resume_priya_patel": {
        "format": "docx",
        "name": "Priya Patel",
        "contact": "priya.patel@example.com | (555) 010-2468 | San Francisco, CA",
        "sections": {
            "Summary": [
                "Senior machine learning engineer with 9 years of Python experience shipping NLP "
                "and large language model systems to production.",
            ],
            "Technical Skills": [
                "Python, PyTorch, TensorFlow, Hugging Face Transformers, LLMs, NLP, "
                "scikit-learn, MLflow, Docker, Kubernetes, AWS SageMaker, SQL",
            ],
            "Work Experience": [
                "Staff ML Engineer, Lumina AI (Jan 2021 – Present)",
                "- Fine-tuned transformer models in PyTorch for document classification at 99.1% "
                "precision.",
                "- Built a retrieval-augmented generation service in Python serving 3M queries "
                "per month.",
                "Machine Learning Engineer, Vertex Analytics (Jun 2017 – Dec 2020)",
                "- Deployed NLP entity-extraction models with TensorFlow Serving on Kubernetes.",
                "- Wrote Python feature pipelines and model monitoring with MLflow.",
            ],
            "Education": [
                "Ph.D. Computer Science (Natural Language Processing), Stanford University, 2017",
                "B.Tech. Computer Science, IIT Bombay, 2012",
            ],
            "Publications": [
                "Efficient Fine-Tuning of Transformers for Low-Resource Domains, ACL 2020.",
            ],
        },
    },
    "resume_daniel_kim": {
        "format": "txt",
        "name": "Daniel Kim",
        "contact": "daniel.kim@example.com | (555) 010-1357 | Denver, CO",
        "sections": {
            "Summary": [
                "Junior Python developer who enjoys building small web apps and automation scripts.",
            ],
            "Skills": [
                "Python, Flask, SQLite, HTML, CSS, JavaScript, Git, pytest",
            ],
            "Experience": [
                "Junior Software Developer, Pinecrest Software (Jul 2024 – Present)",
                "- Maintained internal Flask tools and wrote pytest suites for them.",
                "Software Engineering Intern, Pinecrest Software (Jun 2023 – Aug 2023)",
                "- Automated report generation with Python scripts, saving 5 hours per week.",
            ],
            "Projects": [
                "Recipe Finder – Flask web app with SQLite search and a REST API.",
            ],
            "Education": ["B.S. Computer Science, University of Colorado Boulder, 2024"],
        },
    },
    "resume_fatima_ahmed": {
        "format": "pdf",
        "name": "Fatima Ahmed",
        "contact": "fatima.ahmed@example.com | (555) 010-8642 | Toronto, ON",
        "sections": {
            "Summary": [
                "Site reliability engineer with 7 years of experience running cloud "
                "infrastructure at scale.",
            ],
            "Skills": [
                "Kubernetes, Terraform, AWS, GCP, Docker, Prometheus, Grafana, Python, Go, Bash, "
                "Ansible, CI/CD, Linux",
            ],
            "Experience": [
                "Senior Site Reliability Engineer, Northwind Cloud (2021 - Present)",
                "- Managed 40 Kubernetes clusters with Terraform and Helm across AWS and GCP.",
                "- Cut incident response time by 50% with Prometheus alerting and Python tooling.",
                "DevOps Engineer, Blue Yonder Systems (2018 - 2021)",
                "- Built CI/CD pipelines in GitLab and automated provisioning with Ansible.",
                "Systems Administrator, Coho Hosting (2017 - 2018)",
                "- Administered Linux servers and wrote Bash automation.",
            ],
            "Education": ["B.Eng. Computer Engineering, University of Toronto, 2017"],
            "Certifications": [
                "Certified Kubernetes Administrator (CKA), 2022",
                "AWS Certified Solutions Architect - Professional, 2021",
            ],
        },
    },
    "resume_liam_oconnor": {
        "format": "docx",
        "name": "Liam O'Connor",
        "contact": "liam.oconnor@example.com | (555) 010-9753 | Boston, MA",
        "sections": {
            "Summary": [
                "Data engineer with 6 years of experience designing batch and streaming data "
                "platforms.",
            ],
            "Skills": [
                "Python, Scala, Apache Spark, Airflow, Kafka, Snowflake, dbt, SQL, AWS Glue, "
                "Databricks, Docker",
            ],
            "Professional Experience": [
                "Senior Data Engineer, Fourth Coffee Analytics (Mar 2022 – Present)",
                "- Migrated nightly ETL from on-prem Hadoop to Databricks and Spark, cutting cost "
                "by 35%.",
                "- Built Kafka streaming ingestion processing 1B events per day.",
                "Data Engineer, Wide World Importers (Jan 2019 – Feb 2022)",
                "- Developed Airflow DAGs and dbt models feeding the Snowflake warehouse.",
                "- Wrote Python and Scala Spark jobs for customer data deduplication.",
            ],
            "Education": ["M.S. Data Science, Northeastern University, 2018"],
        },
    },
    "resume_sofia_rossi": {
        "format": "txt",
        "name": "Sofia Rossi",
        "contact": "sofia.rossi@example.com | (555) 010-1122 | Miami, FL",
        "sections": {
            "Summary": [
                "Mobile engineer with 5 years of experience building iOS and Android apps.",
            ],
            "Skills": [
                "Swift, SwiftUI, Kotlin, Jetpack Compose, Flutter, Dart, Firebase, REST APIs, "
                "Git, XCTest",
            ],
            "Experience": [
                "Senior Mobile Engineer, Trey Research (2022 - Present)",
                "- Led the SwiftUI rewrite of a banking app with 2M monthly active users.",
                "Mobile Developer, Proseware Inc. (2020 - 2022)",
                "- Shipped cross-platform Flutter apps and Kotlin Android features.",
            ],
            "Education": ["B.S. Computer Science, Florida International University, 2020"],
        },
    },
    "resume_wei_chen": {
        "format": "pdf",
        "name": "Wei Chen",
        "contact": "wei.chen@example.com | (555) 010-3344 | New York, NY",
        "sections": {
            "Summary": [
                "Full-stack engineer with 4 years of experience shipping SaaS products end to end.",
            ],
            "Skills": [
                "TypeScript, Node.js, Express, React, GraphQL, Python, MongoDB, PostgreSQL, "
                "Docker, AWS",
            ],
            "Experience": [
                "Full Stack Engineer, Graphic Design Institute (2022 - Present)",
                "- Built GraphQL APIs in Node.js and React dashboards for 300 B2B customers.",
                "Software Engineer, Relecloud (2020 - 2022)",
                "- Developed Python and Node.js microservices backed by MongoDB.",
            ],
            "Projects": [
                "OpenInvoice - open-source invoicing app (React, Node.js) with 1.2k GitHub stars.",
            ],
            "Education": ["B.S. Computer Science, Rutgers University, 2020"],
        },
    },
    "resume_olivia_brown": {
        "format": "txt",
        "name": "Olivia Brown",
        "contact": "olivia.brown@example.com | (555) 010-5566 | Pittsburgh, PA",
        "sections": {
            "Summary": [
                "Computer vision engineer applying deep learning to robotics and medical imaging.",
            ],
            "Technical Skills": [
                "Python, TensorFlow, Keras, PyTorch, OpenCV, C++, CUDA, deep learning, "
                "computer vision, NumPy",
            ],
            "Professional Experience": [
                "Computer Vision Engineer, Contoso Robotics (Feb 2022 – Present)",
                "- Trained TensorFlow object-detection models for warehouse robots at 30 FPS.",
                "- Optimized inference with CUDA and C++ kernels.",
                "Research Assistant, Carnegie Mellon University (Sep 2020 – Jan 2022)",
                "- Developed PyTorch segmentation models for MRI scans in Python.",
            ],
            "Education": ["M.S. Robotics, Carnegie Mellon University, 2022"],
        },
    },
    "resume_carlos_mendes": {
        "format": "docx",
        "name": "Carlos Mendes",
        "contact": "carlos.mendes@example.com | (555) 010-7788 | Austin, TX",
        "sections": {
            "Summary": [
                "MLOps engineer with 7 years of experience in Python, taking machine learning "
                "models from notebook to production.",
            ],
            "Skills": [
                "Python, MLflow, Kubeflow, Docker, Kubernetes, scikit-learn, XGBoost, FastAPI, "
                "Airflow, AWS, Terraform",
            ],
            "Experience": [
                "Senior MLOps Engineer, Datum Corp (2022 – Present)",
                "- Built a Kubeflow training platform used by 25 data scientists.",
                "- Served scikit-learn and XGBoost models via FastAPI with canary releases.",
                "Machine Learning Engineer, Southridge Video (2019 – 2022)",
                "- Productionized recommendation models in Python with MLflow tracking.",
                "Backend Developer, Southridge Video (2017 – 2019)",
                "- Wrote Python services and Airflow jobs for video analytics.",
            ],
            "Education": ["B.S. Computer Engineering, Texas A&M University, 2017"],
        },
    },
}


def _write_pdf(path: Path, data: dict) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, data["name"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, data["contact"], new_x="LMARGIN", new_y="NEXT")
    for heading, lines in data["sections"].items():
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 8, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for line in lines:
            pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def _write_docx(path: Path, data: dict) -> None:
    doc = Document()
    doc.add_heading(data["name"], level=0)
    doc.add_paragraph(data["contact"])
    for heading, lines in data["sections"].items():
        doc.add_heading(heading, level=1)
        for line in lines:
            doc.add_paragraph(line)
    doc.save(str(path))


def _write_txt(path: Path, data: dict) -> None:
    out = [data["name"], data["contact"], ""]
    for heading, lines in data["sections"].items():
        out += [heading.upper(), *lines, ""]
    path.write_text("\n".join(out), encoding="utf-8")


WRITERS = {"pdf": _write_pdf, "docx": _write_docx, "txt": _write_txt}


def generate(output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for stem, data in RESUMES.items():
        path = output_dir / f"{stem}.{data['format']}"
        WRITERS[data["format"]](path, data)
        written.append(path)
    return written


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "resumes"
    for p in generate(target):
        print(f"wrote {p}")
