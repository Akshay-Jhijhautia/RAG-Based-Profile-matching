import pytest
from docx import Document
from fpdf import FPDF

import fs_tools


@pytest.fixture
def txt_file(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("Jane Doe\nSkills: Python, SQL\nBuilt python tools.\n", encoding="utf-8")
    return path


@pytest.fixture
def pdf_file(tmp_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "John Smith - Python Developer", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 10, "Experience with Django and C++", new_x="LMARGIN", new_y="NEXT")
    path = tmp_path / "resume.pdf"
    pdf.output(str(path))
    return path


@pytest.fixture
def docx_file(tmp_path):
    doc = Document()
    doc.add_paragraph("Ada Lovelace")
    doc.add_paragraph("Machine learning engineer using PYTHON daily")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "PyTorch"
    path = tmp_path / "resume.docx"
    doc.save(str(path))
    return path


# --- read_file ---------------------------------------------------------------


def test_read_txt(txt_file):
    result = fs_tools.read_file(str(txt_file))
    assert result["success"] is True
    assert "Skills: Python" in result["content"]
    meta = result["metadata"]
    assert meta["name"] == "resume.txt"
    assert meta["extension"] == ".txt"
    assert meta["size_bytes"] == txt_file.stat().st_size
    assert meta["word_count"] == 8
    assert "modified" in meta


def test_read_txt_latin1_fallback(tmp_path):
    path = tmp_path / "latin.txt"
    path.write_bytes("Jos\xe9 Garc\xeda".encode("latin-1"))
    result = fs_tools.read_file(str(path))
    assert result["success"] is True
    assert result["content"] == "José García"


def test_read_pdf(pdf_file):
    result = fs_tools.read_file(str(pdf_file))
    assert result["success"] is True
    assert "Python Developer" in result["content"]
    assert result["metadata"]["page_count"] == 1


def test_read_docx_includes_tables(docx_file):
    result = fs_tools.read_file(str(docx_file))
    assert result["success"] is True
    assert "Ada Lovelace" in result["content"]
    assert "Skill | PyTorch" in result["content"]
    assert "page_count" not in result["metadata"]


def test_read_missing_file(tmp_path):
    result = fs_tools.read_file(str(tmp_path / "nope.pdf"))
    assert result["success"] is False
    assert "not found" in result["error"].lower()


def test_read_directory(tmp_path):
    result = fs_tools.read_file(str(tmp_path))
    assert result["success"] is False
    assert "directory" in result["error"].lower()


def test_read_unsupported_extension(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG")
    result = fs_tools.read_file(str(path))
    assert result["success"] is False
    assert "unsupported" in result["error"].lower()


def test_read_corrupt_pdf(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"this is not a pdf at all")
    result = fs_tools.read_file(str(path))
    assert result["success"] is False
    assert "error" in result


def test_read_corrupt_docx(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    result = fs_tools.read_file(str(path))
    assert result["success"] is False


def test_read_empty_file_warns(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("")
    result = fs_tools.read_file(str(path))
    assert result["success"] is True
    assert "warning" in result


# --- list_files --------------------------------------------------------------


@pytest.fixture
def populated_dir(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"x")
    (tmp_path / "B.PDF").write_bytes(b"xy")
    (tmp_path / "c.txt").write_text("hello")
    (tmp_path / "d.docx").write_bytes(b"z")
    (tmp_path / ".hidden.txt").write_text("secret")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "nested.pdf").write_bytes(b"n")
    return tmp_path


def test_list_all_files(populated_dir):
    names = [f["name"] for f in fs_tools.list_files(str(populated_dir))]
    assert names == ["a.pdf", "B.PDF", "c.txt", "d.docx"]


def test_list_file_metadata(populated_dir):
    entry = next(f for f in fs_tools.list_files(str(populated_dir)) if f["name"] == "B.PDF")
    assert entry["size_bytes"] == 2
    assert entry["extension"] == ".pdf"
    assert entry["path"].endswith("B.PDF")
    assert "modified" in entry


@pytest.mark.parametrize("ext", [".pdf", "pdf", ".PDF", "PDF"])
def test_list_filter_extension_variants(populated_dir, ext):
    names = [f["name"] for f in fs_tools.list_files(str(populated_dir), ext)]
    assert names == ["a.pdf", "B.PDF"]


def test_list_filter_no_matches(populated_dir):
    assert fs_tools.list_files(str(populated_dir), ".csv") == []


def test_list_missing_directory(tmp_path):
    with pytest.raises(FileNotFoundError):
        fs_tools.list_files(str(tmp_path / "missing"))


def test_list_path_is_file(txt_file):
    with pytest.raises(NotADirectoryError):
        fs_tools.list_files(str(txt_file))


# --- write_file --------------------------------------------------------------


def test_write_creates_nested_dirs(tmp_path):
    target = tmp_path / "out" / "deep" / "summary.md"
    result = fs_tools.write_file(str(target), "# Summary\nPython dev")
    assert result["success"] is True
    assert result["created"] is True
    assert result["bytes_written"] == len("# Summary\nPython dev".encode())
    assert target.read_text() == "# Summary\nPython dev"


def test_write_overwrite(tmp_path):
    target = tmp_path / "file.txt"
    target.write_text("old")
    result = fs_tools.write_file(str(target), "new")
    assert result["success"] is True
    assert result["created"] is False
    assert target.read_text() == "new"


def test_write_unicode(tmp_path):
    target = tmp_path / "u.txt"
    result = fs_tools.write_file(str(target), "Résumé — ✓")
    assert result["success"] is True
    assert target.read_text(encoding="utf-8") == "Résumé — ✓"


@pytest.mark.parametrize("name", ["summary.pdf", "summary.docx"])
def test_write_rejects_binary_formats(tmp_path, name):
    result = fs_tools.write_file(str(tmp_path / name), "text")
    assert result["success"] is False
    assert not (tmp_path / name).exists()


def test_write_to_directory_path(tmp_path):
    result = fs_tools.write_file(str(tmp_path), "text")
    assert result["success"] is False


# --- search_in_file ----------------------------------------------------------


def test_search_case_insensitive(txt_file):
    result = fs_tools.search_in_file(str(txt_file), "PYTHON")
    assert result["success"] is True
    assert result["match_count"] == 2
    assert [m["match"] for m in result["matches"]] == ["Python", "python"]
    assert [m["line_number"] for m in result["matches"]] == [2, 3]


def test_search_context(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("a" * 200 + " Kubernetes " + "b" * 200)
    match = fs_tools.search_in_file(str(path), "kubernetes")["matches"][0]
    assert "Kubernetes" in match["context"]
    assert match["context"].startswith("…") and match["context"].endswith("…")
    assert len(match["context"]) < 200


def test_search_no_matches(txt_file):
    result = fs_tools.search_in_file(str(txt_file), "Haskell")
    assert result["success"] is True
    assert result["match_count"] == 0
    assert result["matches"] == []


def test_search_regex_special_chars(pdf_file):
    result = fs_tools.search_in_file(str(pdf_file), "C++")
    assert result["success"] is True
    assert result["match_count"] == 1


def test_search_pdf(pdf_file):
    assert fs_tools.search_in_file(str(pdf_file), "django")["match_count"] == 1


def test_search_docx(docx_file):
    result = fs_tools.search_in_file(str(docx_file), "python")
    assert result["match_count"] == 1
    assert result["matches"][0]["match"] == "PYTHON"


@pytest.mark.parametrize("keyword", ["", "   "])
def test_search_empty_keyword(txt_file, keyword):
    result = fs_tools.search_in_file(str(txt_file), keyword)
    assert result["success"] is False


def test_search_missing_file(tmp_path):
    result = fs_tools.search_in_file(str(tmp_path / "missing.txt"), "python")
    assert result["success"] is False
    assert "not found" in result["error"].lower()
